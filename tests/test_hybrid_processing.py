from __future__ import annotations

import tempfile
import time
import unittest
from pathlib import Path
from zipfile import ZipFile

import fitz
from docx import Document

import app.main as main
from app.document_parsers import (
    DocumentItem,
    analyze_pdf_page,
    extract_docx,
    extract_hwpx,
    extract_metrics,
)


class FakeOcrRunner:
    def __init__(self) -> None:
        self.calls: list[Path] = []

    def run_page(self, image_path: Path, page_number: int, progress_callback=None):
        self.calls.append(image_path)
        if progress_callback:
            progress_callback(0.5, "모의 OCR")
        return [
            DocumentItem(
                page=page_number,
                type="paragraph",
                text=f"OCR:{image_path.stem}",
                markdown=f"OCR:{image_path.stem}",
                section_path=[],
                source_text=f"OCR:{image_path.stem}",
                bbox=[2, 2, 40, 20],
            )
        ]


def make_job(job_id: str, total_pages: int) -> main.JobState:
    now = time.time()
    return main.JobState(
        id=job_id,
        model=main.MODEL_NAME,
        filename="test.pdf",
        title="test",
        document_type="general",
        rag_eligible=True,
        security_level="internal",
        access_groups=[],
        status="running",
        total_pages=total_pages,
        pages=[main.PageStatus(page=index, status="waiting") for index in range(1, total_pages + 1)],
        results=[],
        created_at=now,
        updated_at=now,
    )


def make_png(path: Path, width: int = 320, height: int = 180) -> None:
    pixmap = fitz.Pixmap(fitz.csRGB, fitz.IRect(0, 0, width, height), False)
    pixmap.clear_with(245)
    pixmap.save(path)


class HybridProcessingTests(unittest.TestCase):
    def setUp(self) -> None:
        self.temp_dir = tempfile.TemporaryDirectory()
        self.root = Path(self.temp_dir.name)
        self.previous_result_dir = main.RESULT_DIR
        self.previous_page_dir = main.PAGE_DIR
        main.RESULT_DIR = self.root / "results"
        main.PAGE_DIR = self.root / "pages"

    def tearDown(self) -> None:
        main.RESULT_DIR = self.previous_result_dir
        main.PAGE_DIR = self.previous_page_dir
        self.temp_dir.cleanup()

    def test_pdf_page_analysis_separates_native_and_scanned_pages(self) -> None:
        pdf_path = self.root / "mixed.pdf"
        image_path = self.root / "scan.png"
        make_png(image_path, 600, 800)
        document = fitz.open()
        native_page = document.new_page()
        native_page.insert_text(
            (72, 72),
            "Digital PDF native text layer with enough characters for reliable extraction "
            "and direct indexing without full page OCR.",
        )
        scanned_page = document.new_page()
        scanned_page.insert_image(scanned_page.rect, filename=str(image_path))
        document.save(pdf_path)
        document.close()

        with fitz.open(pdf_path) as reopened:
            # 이미지가 없는 텍스트 페이지는 OCR 없이 원문 추출만 수행한다.
            self.assertEqual(analyze_pdf_page(reopened[0]).mode, "native")
            self.assertEqual(analyze_pdf_page(reopened[1]).mode, "ocr")

    def test_pages_without_images_skip_ocr(self) -> None:
        pdf_path = self.root / "text-only.pdf"
        document = fitz.open()
        text_page = document.new_page()
        text_page.insert_text(
            (72, 72),
            "Text only page with enough characters for reliable native extraction "
            "and therefore no OCR should run on this page at all.",
        )
        document.new_page()  # 빈 페이지
        document.save(pdf_path)
        document.close()

        with fitz.open(pdf_path) as reopened:
            self.assertEqual(analyze_pdf_page(reopened[0]).mode, "native")
            self.assertEqual(analyze_pdf_page(reopened[1]).mode, "native")

        job = make_job("pdf-text-only", 2)
        runner = FakeOcrRunner()
        main.process_pdf(job, pdf_path, runner, self.root / "text-assets", "/assets/pdf-text-only")

        self.assertEqual(len(runner.calls), 0)
        self.assertEqual([page.status for page in job.pages], ["done", "done"])
        self.assertTrue(any("Text only page" in item.text for item in job.results))
        # OCR이 없는 페이지는 2배 고해상도 렌더링을 생략하고 미리보기만 만든다.
        self.assertFalse((main.PAGE_DIR / job.id / "page_0001.png").exists())
        self.assertTrue((self.root / "text-assets" / "page_0001_source.png").exists())

    def test_pdf_processing_uses_region_ocr_and_full_page_ocr(self) -> None:
        pdf_path = self.root / "hybrid.pdf"
        image_path = self.root / "figure.png"
        make_png(image_path)
        document = fitz.open()
        native_page = document.new_page()
        native_page.insert_text(
            (72, 72),
            "Native page text remains authoritative and contains enough characters "
            "to avoid unnecessary full page OCR processing.",
        )
        native_page.insert_image(fitz.Rect(72, 120, 392, 300), filename=str(image_path))
        scan_page = document.new_page()
        scan_page.insert_image(scan_page.rect, filename=str(image_path))
        document.save(pdf_path)
        document.close()

        job = make_job("pdf-hybrid", 2)
        runner = FakeOcrRunner()
        main.process_pdf(
            job,
            pdf_path,
            runner,
            self.root / "assets",
            "/assets/pdf-hybrid",
        )

        self.assertEqual([page.mode for page in job.pages], ["native+region-ocr", "ocr"])
        self.assertEqual(len(runner.calls), 2)
        self.assertTrue(any("region" in path.stem for path in runner.calls))
        self.assertTrue(any(path.stem == "page_0002" for path in runner.calls))
        self.assertTrue(any("Native page text" in item.text for item in job.results))

    def test_docx_native_content_and_embedded_image_are_both_processed(self) -> None:
        image_path = self.root / "embedded.png"
        make_png(image_path)
        docx_path = self.root / "sample.docx"
        document = Document()
        document.add_heading("Native heading", level=1)
        document.add_paragraph("Native DOCX paragraph")
        document.add_picture(str(image_path))
        document.save(docx_path)

        asset_dir = self.root / "docx-assets"
        asset_dir.mkdir()
        items = extract_docx(docx_path, asset_dir, "/assets/docx")
        runner = FakeOcrRunner()
        job = make_job("docx-hybrid", 1)
        enriched = main.enrich_embedded_images_with_ocr(job, items, runner, asset_dir, 35, 92)

        self.assertTrue(any(item.text == "Native DOCX paragraph" for item in enriched))
        self.assertTrue(any(item.type == "image" for item in enriched))
        self.assertTrue(any(item.text.startswith("OCR:") for item in enriched))
        self.assertEqual(len(runner.calls), 1)

    def test_hwpx_xml_content_and_bindata_image_are_both_processed(self) -> None:
        image_path = self.root / "hwpx-image.png"
        make_png(image_path)
        hwpx_path = self.root / "sample.hwpx"
        section_xml = (
            '<?xml version="1.0" encoding="UTF-8"?>'
            '<hs:sec xmlns:hs="urn:hancom:section" xmlns:hp="urn:hancom:paragraph">'
            "<hp:p><hp:run><hp:t>HWPX 원문 문단</hp:t></hp:run></hp:p>"
            "</hs:sec>"
        )
        with ZipFile(hwpx_path, "w") as archive:
            archive.writestr("Contents/section0.xml", section_xml)
            archive.write(image_path, "BinData/image1.png")

        asset_dir = self.root / "hwpx-assets"
        asset_dir.mkdir()
        items = extract_hwpx(hwpx_path, asset_dir, "/assets/hwpx")
        runner = FakeOcrRunner()
        job = make_job("hwpx-hybrid", 1)
        enriched = main.enrich_embedded_images_with_ocr(job, items, runner, asset_dir, 35, 92)

        self.assertTrue(any(item.text == "HWPX 원문 문단" for item in enriched))
        self.assertTrue(any(item.type == "image" for item in enriched))
        self.assertTrue(any(item.text.startswith("OCR:") for item in enriched))
        self.assertEqual(len(runner.calls), 1)

    def test_standalone_image_runs_full_ocr(self) -> None:
        image_path = self.root / "upload.png"
        make_png(image_path)
        asset_dir = self.root / "image-assets"
        asset_dir.mkdir()
        job = make_job("image-ocr", 1)
        runner = FakeOcrRunner()

        main.process_image(job, image_path, runner, asset_dir, "/assets/image-ocr")

        self.assertEqual(job.pages[0].mode, "ocr")
        self.assertEqual(job.pages[0].status, "done")
        self.assertEqual(len(runner.calls), 1)
        self.assertTrue(any(item.type == "page_image" for item in job.results))


def crashing_worker(request_queue, response_queue):  # pragma: no cover - 자식 프로세스에서 실행
    request_queue.get()
    import os

    os._exit(11)


def echo_worker(request_queue, response_queue):  # pragma: no cover - 자식 프로세스에서 실행
    while True:
        request = request_queue.get()
        if request is None:
            return
        request_id, page_path = request
        response_queue.put((request_id, "ok", [{"text": f"echo:{page_path}", "type": "text", "bbox": None}]))


class PaddleWorkerClientTests(unittest.TestCase):
    def test_worker_crash_is_detected_and_worker_restarts(self) -> None:
        client = main.PaddleWorkerClient(worker_target=crashing_worker)
        with self.assertRaises(main.PaddleWorkerCrash):
            client.run("crash.png", timeout_seconds=30)

        # 크래시 후 다음 요청에서 워커가 자동 재기동되어 정상 응답을 받는다.
        client._worker_target = echo_worker
        blocks = client.run("next.png", timeout_seconds=30)
        self.assertEqual(blocks[0]["text"], "echo:next.png")
        client.terminate()


class NumericChartImageTests(unittest.TestCase):
    def test_extract_metrics_keeps_units_and_skips_plain_integers(self) -> None:
        metrics = extract_metrics("취업률 87.5% 달성, 예산 5,000,000원, 2026 기준 ６５％ 참여")
        raws = [metric["raw"] for metric in metrics]
        self.assertIn("87.5%", raws)
        self.assertIn("5,000,000원", raws)
        self.assertIn("65%", raws)
        self.assertNotIn("2026", raws)
        percent = next(metric for metric in metrics if metric["raw"] == "87.5%")
        self.assertEqual(percent["value"], 87.5)
        self.assertEqual(percent["unit"], "%")

    def test_extract_metrics_prefers_longer_unit(self) -> None:
        metrics = extract_metrics("학생 인건비 65만*6개월*5명")
        self.assertIn({"raw": "6개월", "value": 6.0, "unit": "개월"}, metrics)

    def test_decimal_number_is_not_misclassified_as_heading(self) -> None:
        self.assertFalse(main.looks_like_heading("1.5% 증가한 취업률"))
        self.assertTrue(main.looks_like_heading("1. 소요 예산"))
        self.assertTrue(main.looks_like_heading("제1장 사업 개요"))

    def test_chart_block_gets_axis_mapping_and_description(self) -> None:
        blocks = [
            {
                "type": "chart",
                "text": "| 연도 | 취업률(%) |\n| --- | --- |\n| 2024 | 68 |\n| 2025 | 72 |",
                "bbox": [10, 10, 200, 120],
            }
        ]
        items = main.normalize_blocks(blocks, 1)
        chart_items = [item for item in items if item.type == "chart"]
        self.assertEqual(len(chart_items), 1)
        self.assertEqual(chart_items[0].chart["x_axis"], "연도")
        self.assertIn("취업률(%)", chart_items[0].chart["series"])
        self.assertEqual(len(chart_items[0].chart["rows"]), 2)
        self.assertIn("X축", chart_items[0].description)
        self.assertIn("Y축", chart_items[0].description)

    def test_image_block_gets_description_mapping(self) -> None:
        blocks = [{"type": "image", "text": "시스템 구성도 개념 이미지", "bbox": [0, 0, 50, 50]}]
        items = main.normalize_blocks(blocks, 1)
        self.assertEqual(items[0].type, "image")
        self.assertIn("시스템 구성도", items[0].description)

    def test_repeated_label_at_different_positions_is_kept(self) -> None:
        blocks = [
            {"type": "text", "text": "하이브리드", "bbox": [0, 0, 40, 20]},
            {"type": "text", "text": "하이브리드", "bbox": [200, 0, 240, 20]},
            {"type": "text", "text": "하이브리드", "bbox": [200, 0, 240, 20]},
        ]
        items = main.normalize_blocks(blocks, 1)
        self.assertEqual(len([item for item in items if item.text == "하이브리드"]), 2)


if __name__ == "__main__":
    unittest.main()
